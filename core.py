"""Local OCR, bounded page rendering, project persistence and text exports."""
from dataclasses import dataclass, field, asdict
from pathlib import Path
import io
import json
import os
import tempfile
import zipfile

from PIL import Image, ImageOps

MAX_EDGE = 3200


@dataclass
class Page:
    source: str
    number: int
    image: str
    text: str = ""
    regions: list = field(default_factory=list)

    @property
    def title(self):
        return f"{Path(self.source).name} · {self.number}"


def normalize(image):
    image = ImageOps.exif_transpose(image)
    if image.mode in ("RGBA", "LA") or "transparency" in image.info:
        rgba = image.convert("RGBA")
        image = Image.new("RGB", rgba.size, "white")
        image.paste(rgba, mask=rgba.getchannel("A"))
    else:
        image = image.convert("RGB")
    image.thumbnail((MAX_EDGE, MAX_EDGE), Image.Resampling.LANCZOS)
    return image


def render_pages(path):
    """Yield detached images; PDFium stays confined to the worker thread."""
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(str(path)) as pdf:
            for i in range(len(pdf)):
                page = pdf[i]
                try:
                    width, height = page.get_size()
                    scale = min(3, MAX_EDGE / max(width, height))
                    bitmap = page.render(scale=scale)
                    try:
                        image = bitmap.to_pil().convert("RGB")
                    finally:
                        bitmap.close()
                finally:
                    page.close()
                yield i + 1, image
    else:
        with Image.open(path) as source:
            for i in range(getattr(source, "n_frames", 1)):
                source.seek(i)
                yield i + 1, normalize(source.copy())


class OCR:
    def __init__(self):
        import rapidocr
        from rapidocr import RapidOCR
        models = Path(rapidocr.__file__).parent / "models"
        names = {"Det": "PP-OCRv6_det_small.onnx",
                 "Rec": "PP-OCRv6_rec_small.onnx",
                 "Cls": "ch_ppocr_mobile_v2.0_cls_mobile.onnx"}
        for name in names.values():
            if not (models / name).is_file():
                raise RuntimeError(f"Falta un modelo OCR: {name}. Reinstala la aplicación.")
        params = {f"{key}.model_path": str(models / name) for key, name in names.items()}
        params.update({"Global.log_level": "error", "Global.text_score": 0.35,
                       "EngineConfig.onnxruntime.intra_op_num_threads": 4,
                       "EngineConfig.onnxruntime.inter_op_num_threads": 1})
        self.engine = RapidOCR(params=params)

    def recognize(self, image):
        result = self.engine(image)
        regions = []
        if result.txts:
            for box, text, score in zip(result.boxes, result.txts, result.scores):
                regions.append({"box": box.tolist(), "text": text, "score": float(score)})
        return "\n".join(r["text"] for r in regions), regions


def atomic_write(path, writer):
    path = Path(path)
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix=path.suffix)
    os.close(fd)
    try:
        writer(Path(temporary))
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)


def export(pages, path):
    path = Path(path)
    def write(target):
        if path.suffix.lower() == ".docx":
            from docx import Document
            from docx.shared import Pt
            doc = Document()
            doc.styles["Normal"].font.name = "Calibri"
            doc.styles["Normal"].font.size = Pt(11)
            for i, page in enumerate(pages):
                if i:
                    doc.add_page_break()
                doc.add_heading(page.title, level=1)
                for line in page.text.split("\n"):
                    # XML 1.0 cannot contain most ASCII control characters.
                    line = "".join(c for c in line if c in "\t\n\r" or ord(c) >= 32)
                    doc.add_paragraph(line)
            doc.save(target)
        else:
            markdown = path.suffix.lower() == ".md"
            content = "\n\n".join((f"## {p.title}\n\n" if markdown else f"{p.title}\n\n") + p.text for p in pages)
            target.write_text(content + "\n", encoding="utf-8")
    atomic_write(path, write)


def save_project(pages, path):
    def write(target):
        with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
            records = []
            for i, page in enumerate(pages):
                record = asdict(page)
                record["image"] = f"pages/{i}.png"
                archive.write(page.image, record["image"])
                records.append(record)
            archive.writestr("project.json", json.dumps({"version": 1, "pages": records}, ensure_ascii=False))
    atomic_write(path, write)


def load_project(path, cache):
    """Read known members only; never extract user-controlled filesystem paths."""
    import uuid
    pages = []
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 1024 ** 3:
            raise ValueError("El proyecto supera el límite de 1 GB descomprimido.")
        data = json.loads(archive.read("project.json"))
        if data.get("version") != 1:
            raise ValueError("Versión de proyecto no compatible.")
        for record in data["pages"]:
            image_path = Path(cache) / f"{uuid.uuid4().hex}.png"
            with Image.open(io.BytesIO(archive.read(record["image"]))) as image:
                normalize(image).save(image_path)
            pages.append(Page(str(record["source"]), int(record["number"]), str(image_path),
                              str(record["text"]), record.get("regions", [])))
    return pages
